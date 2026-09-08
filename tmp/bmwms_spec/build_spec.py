from __future__ import annotations

from pathlib import Path
from datetime import date
import math

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK


ROOT = Path(__file__).resolve().parents[2]
OUT_DOCX = ROOT / "output" / "documents" / "BMWMS_Dac_ta_chot_Kho_Product_Inbound_Outbound.docx"

BLUE = "1F4D78"
BLUE2 = "2E74B5"
NAVY = "17365D"
INK = "20262E"
MUTED = "5B6573"
LIGHT_BLUE = "E8EEF5"
LIGHTER_BLUE = "F3F6FA"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "D8DEE8"
GREEN = "1F6B4F"
LIGHT_GREEN = "EAF5EF"
GOLD = "7A5A00"
LIGHT_GOLD = "FFF6DD"
RED = "9B1C1C"
LIGHT_RED = "FCECEC"
WHITE = "FFFFFF"


def set_run_font(run, name="Calibri", size=None, color=INK, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    tr_pr.append(cant)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=MID_GRAY, size=5):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        el = borders.find(tag)
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        set_cant_split(row)
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[min(idx, len(widths_dxa)-1)]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def set_keep_with_next(paragraph, keep=True):
    paragraph.paragraph_format.keep_with_next = keep


def add_field(paragraph, instr):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = instr
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr_text, separate, text, end])
    return run


def configure_numbering(doc):
    numbering = doc.part.numbering_part.element
    used_abs = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    used_num = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abs_bullet = max(used_abs or [0]) + 1
    abs_decimal = abs_bullet + 1
    num_bullet = max(used_num or [0]) + 1
    num_decimal = num_bullet + 1

    def add_abstract(abs_id, fmt, texts):
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abs_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "multilevel")
        abstract.append(multi)
        for ilvl, text in enumerate(texts):
            lvl = OxmlElement("w:lvl")
            lvl.set(qn("w:ilvl"), str(ilvl))
            start = OxmlElement("w:start")
            start.set(qn("w:val"), "1")
            lvl.append(start)
            num_fmt = OxmlElement("w:numFmt")
            num_fmt.set(qn("w:val"), fmt)
            lvl.append(num_fmt)
            lvl_text = OxmlElement("w:lvlText")
            lvl_text.set(qn("w:val"), text)
            lvl.append(lvl_text)
            jc = OxmlElement("w:lvlJc")
            jc.set(qn("w:val"), "left")
            lvl.append(jc)
            ppr = OxmlElement("w:pPr")
            tabs = OxmlElement("w:tabs")
            tab = OxmlElement("w:tab")
            tab.set(qn("w:val"), "num")
            tab.set(qn("w:pos"), str(540 + ilvl * 360))
            tabs.append(tab)
            ppr.append(tabs)
            ind = OxmlElement("w:ind")
            ind.set(qn("w:left"), str(540 + ilvl * 360))
            ind.set(qn("w:hanging"), "270")
            ppr.append(ind)
            lvl.append(ppr)
            abstract.append(lvl)
        numbering.append(abstract)

    add_abstract(abs_bullet, "bullet", ["•", "–", "◦"])
    add_abstract(abs_decimal, "decimal", ["%1.", "%2.", "%3."])
    for num_id, abs_id in ((num_bullet, abs_bullet), (num_decimal, abs_decimal)):
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        aid = OxmlElement("w:abstractNumId")
        aid.set(qn("w:val"), str(abs_id))
        num.append(aid)
        numbering.append(num)
    return num_bullet, num_decimal


def apply_num(paragraph, num_id, level=0):
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = ppr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        ppr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    nid = OxmlElement("w:numId")
    nid.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, nid])


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25

for name, size, color, before, after in (
    ("Heading 1", 16, BLUE2, 18, 10),
    ("Heading 2", 13, BLUE2, 14, 7),
    ("Heading 3", 12, BLUE, 10, 5),
):
    st = styles[name]
    st.font.name = "Calibri"
    st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

for style_name, size, color, bold in (
    ("Spec Title", 27, NAVY, True),
    ("Spec Subtitle", 14, MUTED, False),
    ("Spec Kicker", 10, BLUE2, True),
    ("Code Block", 9, INK, False),
    ("Table Text", 9, INK, False),
):
    if style_name not in styles:
        st = styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
    else:
        st = styles[style_name]
    st.font.name = "Calibri" if style_name != "Code Block" else "Consolas"
    st._element.rPr.rFonts.set(qn("w:ascii"), st.font.name)
    st._element.rPr.rFonts.set(qn("w:hAnsi"), st.font.name)
    st._element.rPr.rFonts.set(qn("w:eastAsia"), st.font.name)
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor.from_string(color)
    st.font.bold = bold
    st.paragraph_format.space_before = Pt(0)
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.line_spacing = 1.15

num_bullet, num_decimal = configure_numbering(doc)


def add_body(text="", bold_prefix=None, italic=False):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, italic=italic)
    else:
        r = p.add_run(text)
        set_run_font(r, italic=italic)
    return p


def add_bullet(text, level=0):
    p = doc.add_paragraph(style="Normal")
    apply_num(p, num_bullet, level)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_number(text, level=0):
    p = doc.add_paragraph(style="Normal")
    apply_num(p, num_decimal, level)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_heading(text, level=1, page_break=False):
    if page_break:
        doc.add_page_break()
    p = doc.add_paragraph(text, style=f"Heading {level}")
    return p


def add_callout(label, text, kind="decision"):
    colors = {
        "decision": (GREEN, LIGHT_GREEN),
        "warning": (RED, LIGHT_RED),
        "note": (BLUE, LIGHTER_BLUE),
        "caution": (GOLD, LIGHT_GOLD),
    }
    accent, fill = colors[kind]
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    set_table_borders(table, color=fill, size=2)
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(f"{label}: ")
    set_run_font(r, color=accent, bold=True)
    r2 = p.add_run(text)
    set_run_font(r2, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_code(text):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    set_table_borders(table, color=MID_GRAY, size=4)
    cell = table.cell(0, 0)
    shade_cell(cell, LIGHT_GRAY)
    p = cell.paragraphs[0]
    p.style = styles["Code Block"]
    p.paragraph_format.space_after = Pt(0)
    for idx, line in enumerate(text.splitlines()):
        if idx:
            p.add_run().add_break()
        r = p.add_run(line)
        set_run_font(r, name="Consolas", size=9, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(headers, rows, widths, font_size=8.6, header_fill=LIGHT_BLUE):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, head in enumerate(headers):
        cell = hdr.cells[i]
        shade_cell(cell, header_fill)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(str(head))
        set_run_font(r, size=font_size, color=NAVY, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            r = p.add_run(str(value))
            set_run_font(r, size=font_size, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_caption(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    set_run_font(r, size=9, color=MUTED, italic=True)
    return p


def add_source(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_run_font(r, size=8.5, color=MUTED, italic=True)


# Running header/footer
header = section.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
hp.paragraph_format.space_after = Pt(0)
r = hp.add_run("BMWMS  |  ĐẶC TẢ NGHIỆP VỤ CHỐT")
set_run_font(r, size=8.5, color=MUTED, bold=True)

footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
fp.paragraph_format.space_after = Pt(0)
r = fp.add_run("Bản 1.0  •  Trang ")
set_run_font(r, size=8.5, color=MUTED)
page_run = add_field(fp, "PAGE")
set_run_font(page_run, size=8.5, color=MUTED)


# Cover
p = doc.add_paragraph(style="Spec Kicker")
p.paragraph_format.space_before = Pt(26)
p.paragraph_format.space_after = Pt(10)
r = p.add_run("BUILDING MATERIALS WAREHOUSE MANAGEMENT SYSTEM")
set_run_font(r, size=10, color=BLUE2, bold=True)

p = doc.add_paragraph(style="Spec Title")
p.paragraph_format.space_after = Pt(8)
r = p.add_run("Đặc tả nghiệp vụ chốt và kế hoạch thay đổi hệ thống")
set_run_font(r, size=27, color=NAVY, bold=True)

p = doc.add_paragraph(style="Spec Subtitle")
p.paragraph_format.space_after = Pt(20)
r = p.add_run("Kho • Product • Luồng nhập/xuất • Tự động cấp phát và trừ tồn")
set_run_font(r, size=14, color=MUTED)

meta = [
    ("Phiên bản", "1.0 - Baseline đề xuất để cập nhật đặc tả và triển khai"),
    ("Ngày lập", "25/08/2026"),
    ("Nhánh đối chiếu", "trungdq"),
    ("Phạm vi", "Một kho; quản lý dòng chảy vật lý; không tài chính, hóa đơn hoặc sức chứa chi tiết"),
    ("Đối tượng", "Nhóm phân tích, phát triển, kiểm thử và nghiệm thu BMWMS"),
]
add_table(["Thông tin", "Giá trị"], meta, [1900, 7460], font_size=9.2, header_fill=LIGHT_GRAY)

add_callout(
    "Kết luận xuyên suốt",
    "Không bắt người dùng quản lý mã lô, nhưng hệ thống phải giữ Lớp tồn theo lần nhập (ReceiptStockLayer) để FIFO, LIFO và FEFO có dữ liệu thực thi. Reservation không phải trừ tồn; picking không mặc nhiên là hàng đã rời kho.",
    "decision",
)

add_heading("Cấu trúc tài liệu", 2)
for item in [
    "1. Mục tiêu, phạm vi và thứ tự ưu tiên nguồn yêu cầu",
    "2. Các quyết định nghiệp vụ đã chốt",
    "3. Thiết kế Kho và vị trí lưu trữ",
    "4. Thiết kế Product, Product Group, HSD và lớp tồn",
    "5. Luồng nhập hàng và putaway",
    "6. Luồng xuất hàng và cơ chế tự động trừ tồn",
    "7. Business Rules chuẩn hóa",
    "8. Mô hình dữ liệu đích và kế hoạch migration",
    "9. Kế hoạch thay đổi code và giao diện",
    "10. Ma trận cập nhật tài liệu đặc tả",
    "11. Kiểm thử, nghiệm thu, triển khai và rủi ro",
    "Phụ lục: state machine, dữ liệu trường, thuật ngữ và nguồn tham khảo",
]:
    add_bullet(item)


add_heading("1. Mục tiêu, phạm vi và nguyên tắc ra quyết định", 1)
add_heading("1.1 Mục tiêu", 2)
add_body(
    "Tài liệu này hợp nhất các quyết định đã được nhóm thảo luận về Kho, Product và nghiệp vụ đưa hàng vào/ra khỏi kho. Nó đồng thời là bản phân tích chênh lệch giữa quyết định mới, Report 3.1 Use Case Specification, Business Process, Sequence Diagram, Business Rule và code/database hiện tại trên nhánh trungdq."
)
for text in [
    "Tạo một baseline duy nhất để nhóm không tiếp tục phát triển theo các cách hiểu khác nhau.",
    "Mô tả rõ thời điểm giữ tồn, cấp phát, lấy hàng, chuyển staging và trừ tồn.",
    "Chỉ ra từng thay đổi cần thực hiện trong database, backend, API, UI, test và tài liệu.",
    "Phân biệt phần bắt buộc để đúng tồn kho với phần có thể để ngoài phạm vi đồ án.",
]:
    add_bullet(text)

add_heading("1.2 Phạm vi chốt", 2)
add_table(
    ["Chủ đề", "Trong phạm vi", "Ngoài phạm vi"],
    [
        ("Kho", "Một kho duy nhất; Zone - Rack - Bin; tầng A-Z và vị trí bin; vị trí hệ thống RECEIVING/QUARANTINE/OUTBOUND_STAGING.", "CRUD nhiều kho; sức chứa theo diện tích, trọng lượng, thể tích; tối ưu 3D."),
        ("Product", "Product Group, thuộc tính, một đơn vị tồn chuẩn, QuantityScale, FIFO/LIFO/FEFO, mô tả HSD, expiry thực tế theo lần nhập.", "Tài chính, giá vốn, hóa đơn; serial từng đơn vị; quy đổi UOM phức tạp nếu chưa cấu hình."),
        ("Inbound", "PO mua về và SO khách trả; kiểm nhận; chất lượng; cất hàng vào bin do Staff xác nhận.", "Phiếu nhập bổ sung cho lệnh thiếu; Goods Receipt kế toán; Invoice."),
        ("Outbound", "SO bán ra và PO trả NCC; reservation; pick plan; picking; staging; final issue; lịch sử bất biến.", "Tối ưu tuyến đường bằng tọa độ; vận tải/chuyến xe; tài chính và chứng từ giao nhận kế toán."),
        ("Transfer", "Chỉ điều chuyển nội bộ độc lập giữa hai vị trí trong cùng kho.", "Coi tác nghiệp Inbound/Outbound là Transfer Order."),
    ],
    [1250, 4050, 4060],
)

add_heading("1.3 Thứ tự ưu tiên nguồn yêu cầu", 2)
add_number("Quyết định nghiệp vụ mới đã được nhóm chốt trong quá trình review.")
add_number("Business Rule được chuẩn hóa trong tài liệu này.")
add_number("Report 3.1 UCs và state machine sau khi được cập nhật theo quyết định mới.")
add_number("Business Process và Sequence Diagram sau khi cập nhật.")
add_number("Code/database hiện tại chỉ là hiện trạng để xác định gap, không phải nguồn đúng khi mâu thuẫn với các mục trên.")
add_callout(
    "Nguyên tắc kiểm soát mâu thuẫn",
    "Mọi yêu cầu cũ về nhiều kho, sức chứa, nhập bổ sung, bắt buộc mã lô hoặc ưu tiên vị trí trước FEFO/FIFO/LIFO phải được đánh dấu Superseded và sửa đồng bộ; không xử lý bằng alias UI kéo dài.",
    "warning",
)

add_heading("1.4 Nguồn đã đối chiếu", 2)
for text in [
    "Report 3.1_UCS: UC01-UC09, UC17-UC32, UC48, UC64-UC65 và Entity State Machines.",
    "BMWMS.xlsx - Business Rule.pdf: BR-01 đến BR-20.",
    "BMWMS.xlsx - Nghiệp vụ.pdf và Thiết lập cấu hình cho từng product group.pdf.",
    "BP-01 Procurement and Inbound Receipt; BP-02 Customer Return Inbound; BP-03 Sales Order and Outbound Issue; BP-04 Internal Location Transfer; BP-06 Return Goods to Supplier.",
    "BMWMS_All_Use_Case_Sequence_Diagrams.drawio: đặc biệt UC22, UC23, UC26, UC30 và UC31.",
    "Database_v3.sql, schema patches và code ASP.NET Core hiện tại trên nhánh trungdq.",
]:
    add_bullet(text)


add_heading("2. Các quyết định nghiệp vụ đã chốt", 1, page_break=True)
decisions = [
    ("D-01", "Một kho duy nhất", "Giữ một bản ghi Warehouse làm neo dữ liệu; bỏ Create/Edit/Delete Warehouse khỏi nghiệp vụ người dùng."),
    ("D-02", "Master vị trí", "CRUD chỉ cho Zone, Rack, Bin; thêm thứ tự tầng/bin để cấp phát ổn định; không quản lý sức chứa."),
    ("D-03", "Vị trí hệ thống", "RECEIVING, QUARANTINE và OUTBOUND_STAGING được seed và khóa xóa; không phải kho mới."),
    ("D-04", "Product Group", "RotationMethod là lựa chọn bắt buộc FIFO/LIFO/FEFO; ShelfLifeDescription là text; hai trường nằm ở màn cấu hình nhóm."),
    ("D-05", "Danh sách Product Group", "Không hiển thị cột Luân chuyển và Quản lý lô/HSD; chỉ hiển thị ở chi tiết/cấu hình nhóm."),
    ("D-06", "Không bắt mã lô", "SupplierLotNumber tùy chọn; loại bỏ validation bắt buộc LotNumber và nhãn Số lô khỏi luồng thông thường."),
    ("D-07", "Lớp tồn nội bộ", "Mỗi lần nhận hợp lệ tạo ReceiptStockLayer, tự sinh mã, giữ ReceivedAt và ExpiryDate khi FEFO."),
    ("D-08", "Đơn vị số lượng", "Một UOM tồn chuẩn trên Product; QuantityScale quyết định số nguyên hoặc tối đa số chữ số thập phân."),
    ("D-09", "Inbound", "Chỉ từ PO mua về hoặc SO khách trả; bỏ hoàn toàn chức năng nhập bổ sung cho lệnh thiếu."),
    ("D-10", "Putaway", "Warehouse Staff được assign kiểm nhận, sau đó tự chọn bin thực tế bằng cây Zone/Rack/Bin và xác nhận số lượng đặt tại từng bin."),
    ("D-11", "Outbound", "Chỉ từ SO bán ra hoặc PO trả NCC; picking nằm trong chi tiết Outbound, không tạo Transfer Order tự động."),
    ("D-12", "Reservation", "Draft SO dùng giữ mềm có hạn nếu cần chống chọn trùng; SO xác nhận chuyển sang giữ cứng; reservation không giảm OnHand."),
    ("D-13", "Cấp phát", "RotationMethod là tiêu chí chính; LocationSequence chỉ là tie-breaker trong cùng nhóm ngày/HSD."),
    ("D-14", "Trừ tồn", "Xác nhận pick chuyển số lượng từ bin sang OUTBOUND_STAGING; xác nhận hàng rời kho mới giảm tổng OnHand của kho."),
    ("D-15", "Assign Staff", "Chỉ WAREHOUSE_STAFF đang active và không có nhiệm vụ Inbound/Outbound active khác được assign."),
]
add_table(["ID", "Quyết định", "Nội dung chốt"], decisions, [750, 2200, 6410], font_size=8.8)


add_heading("3. Thiết kế Kho và vị trí lưu trữ", 1, page_break=True)
add_heading("3.1 Mô hình một kho", 2)
add_callout(
    "Chốt",
    "Bỏ CRUD Warehouse ở UI/API nghiệp vụ. Không xóa bảng Warehouses ngay vì hầu hết khóa ngoại đang cần WarehouseID; thay vào đó biến nó thành cấu hình singleton do deployment/seed quản lý.",
    "decision",
)
add_body("Cấu trúc vật lý chuẩn:")
add_code(
    "WAREHOUSE (singleton)\n"
    "├── ZONE\n"
    "│   └── RACK\n"
    "│       └── BIN / STORAGE LOCATION\n"
    "│           ├── LevelCode: A..Z\n"
    "│           ├── LevelSequence: 1..26\n"
    "│           └── BinPosition: 1..n\n"
    "└── SYSTEM LOCATIONS\n"
    "    ├── RECEIVING\n"
    "    ├── QUARANTINE\n"
    "    └── OUTBOUND_STAGING"
)
add_body(
    "Tầng không cần entity CRUD riêng. Nó là thuộc tính cấu trúc của Bin. Cách này giữ đúng cấp bậc Zone - Rack - Bin mà vẫn biểu diễn được tầng A-Z và vị trí 1-n trong từng rack."
)

add_heading("3.2 Dữ liệu vị trí cần có", 2)
add_table(
    ["Field", "Bắt buộc", "Ý nghĩa/Validation"],
    [
        ("ZoneCode, RackCode", "Có", "Mã duy nhất trong phạm vi cha; không thay đổi sau khi đã có giao dịch nếu việc đổi mã làm mất truy vết."),
        ("LocationCode", "Có", "Mã hiển thị duy nhất, ví dụ Z01-R02-A-01."),
        ("LevelCode", "Có với BIN", "A-Z; chỉ dùng hiển thị và tra cứu."),
        ("LevelSequence", "Có với BIN", "Số 1-26; dùng sắp xếp, không sắp bằng chuỗi A-10/A-2."),
        ("BinPosition", "Có với BIN", "Số nguyên dương 1-n."),
        ("LocationSequence", "Có", "Khóa số ổn định gồm Zone/Rack/Level/Bin; allocation sử dụng ASC/DESC."),
        ("LocationType", "Có", "BIN, RECEIVING, QUARANTINE, OUTBOUND_STAGING."),
        ("IsPutawayAllowed", "Có", "Chỉ BIN hoặc vị trí được phép mới nhận putaway."),
        ("IsPickable", "Có", "Chỉ BIN khả dụng mới được chọn làm nguồn pick."),
        ("Status", "Có", "AVAILABLE, OCCUPIED, BLOCKED, INACTIVE; không trộn ACTIVE với AVAILABLE."),
        ("IsSystemLocation", "Có", "Đánh dấu vị trí seed; không được xóa hoặc sửa loại tùy ý."),
    ],
    [2600, 1200, 5560],
)

add_heading("3.3 Quy tắc vị trí", 2)
for text in [
    "Mã và sequence phải duy nhất trong kho; vị trí phải thuộc đúng Rack/Zone của kho singleton.",
    "Không cho chọn BIN BLOCKED/INACTIVE, Rack/Zone inactive hoặc vị trí IsPickable=false trong outbound.",
    "Không cho putaway vào vị trí IsPutawayAllowed=false.",
    "Bin có thể chứa nhiều Product nếu nhóm không đặt hạn chế; nhưng cùng Product khác ReceiptStockLayer phải vẫn tách dữ liệu.",
    "LocationSequence không được tính động từ LocationCode tại thời điểm cấp phát; phải lưu thành số để kết quả ổn định.",
    "Bỏ mọi kiểm tra MaxWeightKg, MaxVolumeM3, AreaSquareMeter trong create/update/putaway vì sức chứa đã out of scope.",
]:
    add_bullet(text)

add_heading("3.4 Tác động code và database", 2)
add_table(
    ["Lớp", "Hiện trạng", "Thay đổi bắt buộc"],
    [
        ("Database", "Warehouses cho phép nhiều dòng; StorageLocations có Area/Weight/Volume nhưng không có Level/Bin sequence.", "Giữ một warehouse seed; thêm singleton guard; bỏ/deprecate capacity; thêm LevelCode, LevelSequence, BinPosition, LocationSequence, IsSystemLocation."),
        ("Backend", "WarehouseService hỗ trợ Create/Update/Delete và đổi kho chính.", "Bỏ mutation endpoint; chỉ còn GetWarehouseOverview. Mọi request nghiệp vụ tự resolve WarehouseID singleton ở server."),
        ("Storage API", "DTO còn AreaSquareMeter, MaxWeightKg, MaxVolumeM3 và status không đồng nhất.", "Loại capacity field; chuẩn hóa enum/status; thêm fields thứ tự và validation hierarchy."),
        ("UI", "Có trang Warehouse Create/Edit/Delete; Location hiển thị capacity.", "Ẩn/xóa CRUD Warehouse; biến Warehouse thành trang tổng quan; Location CRUD theo cây Zone/Rack/Bin, không có capacity."),
        ("Docs", "UC01/BR-17 còn khái niệm warehouse status; BR-04 bắt capacity.", "UC01 chỉ View singleton; retire BR-04; BR-17 đổi thành kiểm tra cấu hình singleton; cập nhật UC02/ERD/BP."),
    ],
    [1300, 3600, 4460],
)


add_heading("4. Thiết kế Product, HSD và lớp tồn", 1, page_break=True)
add_heading("4.1 Product Group là nơi cấu hình chính sách", 2)
add_body(
    "RotationMethod và ShelfLifeDescription là chính sách nghiệp vụ, không nên cài như ProductAttribute động thông thường. Chúng phải là cột có kiểu dữ liệu rõ trên ProductGroups, nhưng được đặt trong cùng màn hình cấu hình thuộc tính của nhóm."
)
add_table(
    ["Field", "Kiểu", "Rule"],
    [
        ("RotationMethod", "Selection", "Bắt buộc; chỉ FIFO, LIFO hoặc FEFO."),
        ("ShelfLifeDescription", "Text", "Tùy chọn; mô tả như '12 tháng từ NSX' hoặc 'Theo HSD trên bao bì'; không dùng trực tiếp để sort FEFO."),
        ("Status", "Selection", "ACTIVE/INACTIVE; group inactive không được chọn cho Product mới."),
        ("Dynamic Attributes", "Collection", "Các thuộc tính nhận diện sản phẩm: mác, kích thước, màu, đường kính..."),
    ],
    [2700, 1500, 5160],
)
add_callout(
    "Ràng buộc FEFO",
    "ShelfLifeDescription là text chỉ để mô tả. Allocation FEFO bắt buộc dùng ExpiryDate cụ thể trên ReceiptStockLayer của lần nhập. Không được parse text HSD để tự suy diễn ngày hết hạn.",
    "caution",
)

add_heading("4.2 Cách Product nhận chính sách", 2)
for text in [
    "Product thuộc đúng một Product Group và dùng RotationMethod hiệu lực từ Group.",
    "Không duy trì RotationMethod đồng thời ở Product và ProductGroup vì sẽ phát sinh hai nguồn đúng khác nhau.",
    "Nếu một sản phẩm cần phương pháp khác các sản phẩm cùng nhóm, chuyển nó sang nhóm phù hợp; Product-level override để ngoài scope hiện tại.",
    "Danh sách Product Group không thêm cột Rotation/HSD; chi tiết và màn cấu hình phải hiển thị để người quản lý kiểm tra.",
]:
    add_bullet(text)

add_heading("4.3 Đơn vị tính và QuantityScale", 2)
add_table(
    ["Ví dụ UOM", "QuantityScale", "Giá trị hợp lệ", "Giá trị không hợp lệ"],
    [
        ("bao, cây, cuộn", "0", "1; 25; 100", "1.5; 2.25"),
        ("tấn, kg, m², m³", "3", "0.125; 12.500", "0.1234"),
        ("đơn vị cần 2 số lẻ", "2", "1.25; 10.00", "1.257"),
    ],
    [2100, 1500, 2700, 3060],
)
for text in [
    "Tất cả Ordered, Expected, Received, Putaway, Reserved, Picked và Issued Quantity phải dùng cùng QuantityScale của Product UOM.",
    "Database dùng DECIMAL(18,4); application validation giới hạn số chữ số theo QuantityScale.",
    "Lô không phải UOM hoặc quy cách đóng gói. Nếu cần 1 pallet = 40 bao, đó là UOM conversion riêng, chưa nằm trong baseline này.",
]:
    add_bullet(text)

add_heading("4.4 Mã lô và ReceiptStockLayer", 2)
add_body(
    "Mã lô nhà sản xuất dùng cho truy xuất một nhóm hàng sản xuất cùng điều kiện. Hệ thống BMWMS hiện không cần bắt buộc khả năng thu hồi theo lô. Tuy nhiên FIFO/LIFO/FEFO vẫn cần biết số lượng nào đã vào kho khi nào và có HSD nào. Vì thế phải tách khái niệm bên ngoài và nội bộ:"
)
add_table(
    ["Khái niệm", "Người nhập", "Bắt buộc", "Mục đích"],
    [
        ("SupplierLotNumber", "Staff nếu có trên nhãn", "Không", "Truy xuất lô NCC/NSX; để mở cho tương lai."),
        ("ReceiptStockLayerId", "Hệ thống tự sinh", "Có", "Phân biệt từng đợt nhận để tính tuổi tồn, HSD và số lượng theo bin."),
        ("SystemLayerCode", "Hệ thống tự sinh", "Có", "Mã kỹ thuật để audit, ví dụ SL-20260825-000001; không bắt Staff hiểu."),
        ("ReceivedAt", "Hệ thống", "Có", "Khóa chính cho FIFO/LIFO."),
        ("ExpiryDate", "Staff nhập/xác nhận", "Bắt buộc với FEFO", "Khóa chính cho FEFO; phải là ngày hợp lệ."),
    ],
    [2200, 1900, 1500, 3760],
)
add_code(
    "ReceiptStockLayer SL-1001\n"
    "Product: Xi măng PCB40\n"
    "ReceivedAt: 25/08/2026 09:15\n"
    "ExpiryDate: 25/10/2026\n"
    "SupplierLotNumber: null\n"
    "Balances:\n"
    "  Z01-R01-A-01 = 200 bao\n"
    "  Z01-R01-A-02 = 200 bao\n"
    "  Z01-R01-A-03 = 100 bao"
)
add_callout(
    "Lỗi hiện tại phải sửa",
    "Code dùng NO-LOT-{ProductId} cho mọi lần nhập không quản lý lô. Kết quả là các lần nhập của cùng Product bị gộp vào một ProductLot và FirstReceivedDate không còn đại diện cho phần tồn thực tế; FIFO/LIFO sẽ sai.",
    "warning",
)

add_heading("4.5 Tác động code/database/tài liệu", 2)
add_table(
    ["Thành phần", "Thay đổi"],
    [
        ("ProductGroups", "Thêm RotationMethod NOT NULL và ShelfLifeDescription nullable; CHECK FIFO/LIFO/FEFO."),
        ("Products", "Bỏ/deprecate RotationMethod, TrackLot, TrackExpiry, DefaultShelfLifeDays; đọc policy từ ProductGroup."),
        ("ProductLots", "Đổi tên/migration sang ReceiptStockLayers; LotNumber thành SupplierLotNumber nullable; thêm SystemLayerCode và ReceivedAt."),
        ("Product UI", "Không còn checkbox Quản lý lô/HSD; hiển thị chính sách hiệu lực đọc từ Group; HSD mô tả read-only."),
        ("Inbound UI", "Không bắt số lô. FEFO bắt ExpiryDate; SupplierLotNumber chỉ là trường mở rộng tùy chọn."),
        ("Reports", "Đổi Expiring Lot thành Cảnh báo tồn sắp hết hạn; hiển thị Product, ExpiryDate, Receipt reference, Bin và quantity."),
        ("Report 3.1", "UC03/05/07/09/20/22/23/28/30/31/32/48 và state machine Product Lot phải đổi sang ReceiptStockLayer."),
    ],
    [2500, 6860],
    font_size=7.8,
)


add_heading("5. Luồng nhập hàng và putaway", 1, page_break=True)
add_heading("5.1 Các trường hợp nguồn", 2)
add_table(
    ["SourceType", "Tham chiếu", "Ý nghĩa", "Kết quả"],
    [
        ("PURCHASE_ORDER", "PO đã CONFIRMED/PARTIALLY_RECEIVED", "Hàng mua về từ NCC.", "Inbound riêng; kiểm nhận; putaway vào bin."),
        ("SALES_RETURN", "SO đã phát sinh xuất", "Hàng khách trả lại.", "Inbound riêng; bắt buộc kiểm định; tạo lớp tồn mới."),
    ],
    [1900, 2800, 2300, 2360],
)
add_callout(
    "Loại bỏ",
    "Không còn ParentInboundOrder, Supplemental flag hay màn 'Nhập bổ sung cho lệnh thiếu'. Thiếu/hỏng chỉ được ghi nhận và thông báo; nếu cần mua/giao bù thì lập nghiệp vụ PO/Inbound mới hợp lệ.",
    "decision",
)

add_heading("5.2 State machine đích", 2)
add_code(
    "DRAFT\n"
    "  └─ submit + assign WAREHOUSE_STAFF ─> READY\n"
    "READY\n"
    "  └─ Staff bắt đầu kiểm nhận ────────> RECEIVING\n"
    "RECEIVING\n"
    "  └─ chốt đủ mọi dòng/ngoại lệ ─────> RECEIVED\n"
    "RECEIVED\n"
    "  └─ cất đủ hàng đạt vào bin ───────> PUTAWAY_COMPLETED\n"
    "DRAFT hoặc READY ── hủy ────────────> CANCELLED"
)
add_body(
    "Assign là một hành động nghiệp vụ: phiếu chỉ chuyển READY khi đã có Warehouse Staff hợp lệ. Staff bắt đầu kiểm nhận chuyển RECEIVING. Trạng thái RECEIVED chỉ thể hiện đã chốt kiểm nhận, chưa khẳng định hàng đạt đã nằm tại Bin."
)

add_heading("5.3 Kiểm nhận", 2)
add_number("Staff mở đúng phiếu được assign; hệ thống khóa version của phiếu.")
add_number("Mỗi dòng nhập DeliveredQuantity, AcceptedQuantity, DamagedQuantity và ConditionNotes.")
add_number("Validation chạy trên toàn bộ dòng trước khi hoàn tất; không cho bấm xác nhận khi còn dòng chưa xử lý hoặc tổng Accepted + Damaged vượt Delivered.")
add_number("FEFO yêu cầu ExpiryDate; FIFO/LIFO không yêu cầu mã lô nhưng hệ thống tự tạo ReceiptStockLayer.")
add_number("Hàng đạt vào RECEIVING với trạng thái chưa khả dụng; hàng hỏng/chờ xử lý vào QUARANTINE.")
add_number("Thiếu/hỏng được lưu lịch sử và thông báo Purchasing/Warehouse Manager; không sinh inbound bổ sung tự động.")

add_heading("5.4 Putaway", 2)
for text in [
    "Staff thấy cây Zone > Rack > Bin, tìm kiếm theo mã/tên, lọc khu/rack/tầng/trạng thái và xem Product hiện có trong Bin.",
    "Hệ thống ưu tiên vị trí cố định/eligible zone nếu đã cấu hình nhưng Staff là người xác nhận Bin thực tế.",
    "Một ReceiptStockLayer có thể chia vào nhiều Bin; tổng putaway không vượt accepted quantity còn chờ.",
    "Không validate capacity. Nếu vị trí không phù hợp vật lý, Staff chọn vị trí khác và hệ thống lưu người/thời gian/lý do nếu khác đề xuất.",
    "Putaway là thao tác trong phiếu Inbound; không tạo TransferOrder.",
]:
    add_bullet(text)

add_heading("5.5 Cách ghi tồn đề xuất", 2)
add_body("Phương án khuyến nghị phản ánh đủ dòng chảy vật lý:")
add_code(
    "Confirm receipt:\n"
    "  RECEIVING balance +Accepted      (không pickable, không available)\n"
    "  QUARANTINE balance +Damaged/Hold  (không available)\n\n"
    "Confirm putaway:\n"
    "  RECEIVING balance -PutawayQuantity\n"
    "  BIN balance       +PutawayQuantity\n"
    "  ReceiptStockLayer giữ nguyên\n\n"
    "Available sales stock:\n"
    "  Chỉ BIN IsPickable + Layer AVAILABLE - ActiveReservation"
)
add_body(
    "Nếu nhóm muốn tối thiểu hóa transaction, có thể không tạo balance tại RECEIVING và chỉ post tồn khi putaway. Tuy nhiên phương án đó làm hệ thống không phản ánh hàng đang có vật lý tại khu nhận; vì vậy không được khuyến nghị cho baseline cuối."
)

add_heading("5.6 Trường hợp khách trả lại", 2)
for text in [
    "Chỉ cho chọn SO đã có issued quantity và không vượt số đã xuất trừ số đã trả trước đó.",
    "Tạo ReceiptStockLayer mới; không nhập thẳng vào lớp tồn cũ.",
    "Mặc định vào QUARANTINE/PENDING_INSPECTION cho đến khi xác nhận usable condition.",
    "Nếu FEFO, Staff phải xác nhận HSD còn lại; hàng hết hạn không vào AVAILABLE.",
    "Giữ OriginalOutboundDetailId/OriginalReceiptStockLayerId nếu truy ra được để hỗ trợ lịch sử.",
]:
    add_bullet(text)

add_heading("5.7 Gap hiện tại", 2)
add_table(
    ["Gap", "Ảnh hưởng", "Hướng sửa"],
    [
        ("Status DB ASSIGNED/IN_PROGRESS/COMPLETED được map thành READY/RECEIVING/RECEIVED/PUTAWAY_COMPLETED bằng suy luận transaction.", "Khó truy vấn, state machine không rõ và dễ lệch API/UI.", "Dùng status canonical trực tiếp hoặc thêm ExecutionStage rõ; migration dữ liệu."),
        ("ProductLot bắt buộc; NO-LOT dùng chung theo Product.", "Mất tuổi tồn từng lần nhập.", "Mỗi receipt tạo ReceiptStockLayer tự sinh."),
        ("Supplemental fields/parent inbound vẫn tồn tại.", "Trái quyết định bỏ nhập bổ sung.", "Bỏ UI, service, DTO và các cột ParentInboundOrderID/Supplemental."),
        ("InboundOrderDetail đồng thời dùng cho receipt staging và putaway actual.", "Khó phân biệt kiểm nhận với vị trí đã cất.", "Tách InboundReceiptLines và InboundPutawayLines hoặc bổ sung execution type rõ."),
        ("Current putaway đã cho Staff chọn cây Zone/Rack/Bin.", "Phù hợp hướng mới.", "Giữ UX này; bỏ capacity và thay ProductLot bằng ReceiptStockLayer."),
    ],
    [3000, 3000, 3360],
)


add_heading("6. Luồng xuất và tự động trừ tồn", 1, page_break=True)
add_heading("6.1 Nguyên tắc nền", 2)
add_callout(
    "Quy tắc vàng",
    "Tự động hóa có nghĩa là hệ thống tự chọn nguồn tồn, giữ số lượng và sinh kế hoạch lấy. Hệ thống không được tự giảm tồn trước khi Warehouse Staff xác nhận hành động vật lý.",
    "decision",
)
add_table(
    ["Khái niệm", "Thay đổi quantity", "Ý nghĩa"],
    [
        ("OnHand", "Chỉ đổi khi nhận/putaway/pick/stage/issue/adjust được post", "Số lượng vật lý đang nằm tại một location."),
        ("Reserved", "Tăng khi giữ hàng; giảm khi release/consume", "Phần OnHand không được đơn khác dùng."),
        ("Available", "OnHand - active reserved, sau khi lọc trạng thái", "Số có thể cấp phát mới."),
        ("PlannedPick", "Không đổi balance", "Số hệ thống đề xuất Staff lấy tại từng bin/layer."),
        ("Picked/Staged", "Bin giảm, Staging tăng", "Hàng đã rời bin nhưng vẫn còn trong kho."),
        ("Issued", "Staging giảm; tổng kho giảm", "Hàng đã giao ra ngoài phạm vi kho."),
    ],
    [1800, 3000, 4560],
)

add_heading("6.2 Nguồn Outbound", 2)
add_table(
    ["SourceType", "Nguồn", "Cấp phát"],
    [
        ("SALES_ORDER", "SO đã CONFIRMED/PARTIALLY_FULFILLED", "Dùng hard reservation của SO; tạo PickPlan theo reservation."),
        ("PURCHASE_RETURN", "PO đã có hàng nhận và còn lượng được trả", "Tạo reservation khi Outbound return được duyệt; ưu tiên lớp tồn có nguồn từ chính PO đó."),
    ],
    [2000, 3300, 4060],
)

add_heading("6.3 Reservation hai lớp", 2)
add_heading("6.3.1 SO Draft - giữ mềm", 3)
for text in [
    "Mục đích: tránh hai Sales cùng chọn một lượng hàng trong lúc soạn SO.",
    "Có ExpiresAt, ví dụ 30 phút; thao tác sửa/lưu Draft hợp lệ có thể gia hạn.",
    "Hết hạn, đóng tab lâu hoặc hủy Draft thì chuyển EXPIRED/RELEASED và trả lại Available.",
    "Không giảm OnHand và không được tồn tại vĩnh viễn.",
]:
    add_bullet(text)

add_heading("6.3.2 SO Confirmed - giữ cứng", 3)
for text in [
    "Khi xác nhận SO, hệ thống recheck và atomically chuyển soft hold thành hard reservation hoặc cấp phát lại.",
    "Reservation chỉ rõ SalesOrderDetail + Product + ReceiptStockLayer + StorageLocation + Quantity.",
    "Hủy Outbound nhưng SO còn hiệu lực thì giữ hard reservation để có thể lập lại Outbound; chỉ SO cancel/reduce hoặc quyết định release mới giải phóng.",
    "Cách này thay thế quy tắc cũ UC29 luôn release reservation khi hủy Outbound.",
]:
    add_bullet(text)

add_heading("6.4 Bộ lọc tồn đủ điều kiện", 2)
add_body("Allocation Engine chỉ xét dòng thỏa đồng thời:")
for text in [
    "Product đúng và ACTIVE; ProductGroup ACTIVE; UOM/QuantityScale hợp lệ.",
    "Location thuộc kho singleton, LocationType=BIN, IsPickable=true, status AVAILABLE/OCCUPIED; Zone và Rack active.",
    "ReceiptStockLayer status AVAILABLE; không EXPIRED, BLOCKED, QUARANTINED hoặc DEPLETED.",
    "OnHandQuantity - ReservedQuantity > 0.",
    "Không bị khóa kiểm kê hoặc bởi tác nghiệp đồng thời.",
    "FEFO có ExpiryDate hợp lệ và chưa hết hạn theo ngày nghiệp vụ.",
    "Không sử dụng reservation thuộc SO/Outbound khác.",
]:
    add_bullet(text)

add_heading("6.5 Thứ tự cấp phát theo phương pháp", 2)
add_table(
    ["Method", "Khóa chính", "Tie-break cùng ngày/HSD", "Vị trí"],
    [
        ("FIFO", "ReceivedDate ASC", "ReceivedAt ASC hoặc SystemLayerCode ASC", "LocationSequence ASC: tầng A-Z, bin nhỏ-lớn."),
        ("LIFO", "ReceivedDate DESC", "ReceivedAt DESC hoặc SystemLayerCode DESC", "LocationSequence DESC: tầng Z-A, bin lớn-nhỏ."),
        ("FEFO", "ExpiryDate ASC", "ReceivedDate ASC, ReceivedAt ASC", "LocationSequence ASC trong cùng HSD/ngày nhập."),
    ],
    [1200, 2300, 3000, 2860],
)
add_callout(
    "Độ ưu tiên",
    "Rotation key luôn đứng trước LocationSequence. 'Gần cửa xuất' hoặc mã bin chỉ được dùng sau khi các candidate có cùng mức ưu tiên FIFO/LIFO/FEFO. BR-07 cũ phải được sửa theo nguyên tắc này.",
    "caution",
)

add_heading("6.6 Thuật toán cấp phát", 2)
add_code(
    "INPUT: productId, requestedQty, sourceReference, rotationMethod\n\n"
    "1. candidates = eligible InventoryBalance rows\n"
    "2. sort candidates:\n"
    "   FIFO: ReceivedDate ASC, ReceivedAt ASC, LocationSequence ASC, LayerId ASC\n"
    "   LIFO: ReceivedDate DESC, ReceivedAt DESC, LocationSequence DESC, LayerId DESC\n"
    "   FEFO: ExpiryDate ASC, ReceivedDate ASC, ReceivedAt ASC,\n"
    "         LocationSequence ASC, LayerId ASC\n"
    "3. remaining = requestedQty\n"
    "4. for each candidate while remaining > 0:\n"
    "      available = OnHand - Reserved\n"
    "      allocated = min(remaining, available)\n"
    "      create/update Reservation(candidate, allocated)\n"
    "      create PickPlanLine(candidate, allocated, sequenceNo)\n"
    "      remaining -= allocated\n"
    "5. if remaining > 0: rollback entire hard allocation\n"
    "6. commit with audit + idempotency key"
)
add_body(
    "Thuật toán greedy theo thứ tự trên vừa giữ đúng rotation vừa có xu hướng dùng hết candidate đầu trước khi sang candidate khác, nhờ đó tránh chia nhỏ qua quá nhiều bin. Không cần một bộ tối ưu tuyến đường phức tạp trong scope hiện tại."
)

add_heading("6.7 Pick Plan và màn Staff", 2)
add_body("Sau allocation, hệ thống sinh PickPlanLine chứ không bắt Staff tự tìm tồn:")
add_table(
    ["Sequence", "Product", "Nguồn", "Lớp tồn/HSD", "Phải lấy"],
    [
        ("1", "XM001 - Xi măng PCB40", "Z01/R01/A-01", "Nhập 20/08; HSD -", "50 bao"),
        ("2", "XM001 - Xi măng PCB40", "Z01/R01/A-02", "Nhập 20/08; HSD -", "30 bao"),
        ("3", "S001 - Sơn ngoại thất", "Z02/R03/B-09", "HSD 10/09/2026", "12 thùng"),
    ],
    [900, 2350, 1850, 2600, 1660],
)
add_body("Màn Process Outbound phải hiển thị:")
for text in [
    "Nguồn SO/PO, khách hàng/NCC, yêu cầu, đã giữ, đã pick, đã issue và còn lại.",
    "RotationMethod hiệu lực và giải thích ngắn về thứ tự.",
    "Cây/mã Zone - Rack - Level - Bin, ReceivedDate và ExpiryDate nếu FEFO.",
    "PlannedQuantity, ActualPickedQuantity, variance và exception reason.",
    "Các nút: Bắt đầu lấy; Xác nhận tại vị trí; Báo thiếu; Cấp phát lại; Hoàn tất lấy; Xác nhận xuất khỏi kho.",
]:
    add_bullet(text)

add_heading("6.8 Picking, staging và final issue", 2)
add_code(
    "READY\n"
    "  └─ Staff bắt đầu ───────────────> PICKING\n"
    "PICKING\n"
    "  ├─ xác nhận từng PickPlanLine\n"
    "  │    BIN OnHand -= actualPicked\n"
    "  │    OUTBOUND_STAGING OnHand += actualPicked\n"
    "  │    reservation consumed tương ứng\n"
    "  └─ đủ mọi dòng ────────────────> PICKED\n"
    "PICKED\n"
    "  └─ xác nhận hàng rời kho\n"
    "       OUTBOUND_STAGING OnHand -= issued\n"
    "       Warehouse total OnHand giảm\n"
    "       SO fulfilled tăng ─────────> ISSUED"
)
add_body(
    "Nếu Report 3.1 phải giữ nguyên số lượng UC, UC31 được tách thành UC31.1 Confirm Pick và UC31.2 Confirm Final Issue. Nếu cho phép thêm UC, nên tạo UC66 Confirm Final Issue để state machine và sequence diagram rõ hơn."
)

add_heading("6.9 Short pick và cấp phát lại", 2)
add_number("Staff nhập ActualPicked nhỏ hơn Planned và chọn lý do: thiếu vật lý, vị trí bị khóa, hàng hỏng, không tìm thấy.")
add_number("Hệ thống chỉ post số thực lấy; không tự điều chỉnh phần chênh lệch khỏi sổ tồn.")
add_number("Phần reservation chưa dùng tại candidate cũ được release hoặc giữ khóa điều tra theo quyết định.")
add_number("Allocation Engine chạy lại cho phần còn thiếu trên candidate kế tiếp.")
add_number("Nếu không đủ, Outbound giữ PICKING với exception hoặc chuyển NEEDS_REVIEW; Warehouse Manager quyết định xuất thiếu/hủy phần còn lại.")

add_heading("6.10 Override rotation/location", 2)
for text in [
    "Staff không được tùy ý chọn vị trí nằm ngoài PickPlan.",
    "Nếu vị trí đề xuất không dùng được, yêu cầu Reallocate; hệ thống giữ/release lại atomically.",
    "Chọn một lớp FIFO muộn hơn, LIFO cũ hơn hoặc FEFO có HSD xa hơn trong khi còn candidate ưu tiên phải cần Warehouse Manager authorization và reason.",
    "Override được ghi audit với proposed source, actual source, actor, approver và timestamp.",
]:
    add_bullet(text)

add_heading("6.11 Hủy và đảo giao dịch", 2)
add_table(
    ["Thời điểm", "Cho phép", "Cách xử lý"],
    [
        ("DRAFT/READY, chưa pick", "Hủy Outbound", "Hủy PickPlan. Hard reservation của SO giữ nếu SO còn hiệu lực; reservation PO return được release."),
        ("PICKING, đã pick một phần", "Không hủy trực tiếp", "Dừng task; trả hàng từ staging về bin bằng compensating movement; sau đó mới cancel/review."),
        ("PICKED", "Không hủy trực tiếp", "Return-to-bin workflow từ staging, giữ audit."),
        ("ISSUED", "Không hủy", "Tạo Customer Return Inbound hoặc nghiệp vụ hoàn trả phù hợp."),
    ],
    [2000, 1900, 5460],
)

add_heading("6.12 Đồng thời, idempotency và tính nguyên tử", 2)
for text in [
    "Allocation/reservation chạy trong transaction; khóa hàng theo Inventory row hoặc dùng RowVersion để chống hai SO lấy cùng quantity.",
    "Mọi command POST xác nhận phải có IdempotencyKey; nhấn đúp/retry không tạo hai transaction.",
    "Inventory update, reservation status, PickPlan/PickExecution, SO/Outbound status và audit phải commit/rollback cùng nhau.",
    "InventoryTransaction bất biến; sửa sai bằng compensating transaction, không UPDATE/DELETE lịch sử.",
    "Mọi sort phải deterministic với LayerId/InventoryId cuối cùng để cùng dữ liệu luôn ra cùng PickPlan.",
]:
    add_bullet(text)

add_heading("6.13 Gap code hiện tại và refactor cần làm", 2)
add_table(
    ["Hiện trạng", "Vấn đề", "Thay đổi"],
    [
        ("InventoryRepository.ReserveStockForOrderAsync chỉ tách FEFO và else.", "LIFO rơi vào FIFO.", "Strategy pattern/IAllocationPolicy với FIFO, LIFO, FEFO riêng."),
        ("BuildInventoryAllocationsAsync OrderBy LocationCode trước ExpiryDate.", "FEFO có thể lấy vị trí trước HSD.", "Rotation key trước, LocationSequence tie-break sau."),
        ("BuildSalesAllocationsAsync order theo LocationCode.", "Pick route có thể đảo reservation priority.", "PickPlan dùng allocation sequence đã lưu, không sort lại bằng mã."),
        ("ExecutePickAsync tạo OUTBOUND delta âm và hoàn tất khi đủ.", "Picking bị đồng nhất với final issue.", "Tách ConfirmPick và ConfirmIssue; thêm OUTBOUND_STAGING."),
        ("ProductLotId bắt buộc trên Inventory/Reservation/Detail.", "Giao diện buộc khái niệm lô và gộp NO-LOT.", "Chuyển FK sang ReceiptStockLayerId; SupplierLotNumber nullable."),
        ("TakeAllocations trả toàn bộ candidate rows.", "Không biểu diễn exact planned quantity dòng cuối.", "Sinh PickPlanLine với allocated=min(remaining, available)."),
        ("Hủy Outbound giữ reservation SO theo comment; UC29 cũ yêu cầu release.", "Tài liệu/code mâu thuẫn.", "Chốt reservation thuộc vòng đời SO; cập nhật UC29 và chỉ release khi SO giảm/hủy."),
        ("Status ASSIGNED/IN_PROGRESS/COMPLETED được normalize thành READY/ISSUING/ISSUED.", "DB/API/UI không cùng ngôn ngữ.", "Migration sang status canonical và state transition service duy nhất."),
    ],
    [3200, 2800, 3360],
)


add_heading("7. Business Rules chuẩn hóa", 1, page_break=True)
add_heading("7.1 Đối chiếu BR hiện có", 2)
br_rows = [
    ("BR-01 FIFO", "Giữ, sửa", "Ưu tiên ReceiptStockLayer có ReceivedDate sớm; không dùng từ 'lô' bắt buộc."),
    ("BR-02 FEFO", "Giữ, sửa", "Ưu tiên ExpiryDate sớm; SupplierLotNumber không bắt buộc nhưng ExpiryDate bắt buộc."),
    ("BR-03 Fixed Location", "Giữ, làm mềm", "Fixed/eligible location là đề xuất và ràng buộc khu; Staff xác nhận bin thực tế."),
    ("BR-04 Capacity", "Retire", "Không quản lý sức chứa; xóa khỏi validation, UI và tài liệu."),
    ("BR-05 Available Stock", "Giữ", "Cấp phát không vượt Available; lọc status và reservation."),
    ("BR-06 Reserved Exclusion", "Giữ", "Reservation của đơn khác không được cấp phát lại."),
    ("BR-07 Nearest Location", "Sửa", "Thay bằng LocationSequence tie-break sau rotation; không đứng trước FEFO/FIFO/LIFO."),
    ("BR-08 Single Order Picking", "Giữ", "Mỗi task xử lý một Outbound; một Staff không có hai task active."),
    ("BR-09 Auto Picking Task", "Giữ, sửa", "Sinh PickPlan trong Outbound, không sinh TransferOrder."),
    ("BR-10 Inventory Update", "Làm rõ", "Update theo posting point: receipt/putaway, pick-to-staging, final issue."),
    ("BR-11 Transaction Logging", "Giữ", "Ledger bất biến; corrections là compensating transactions."),
    ("BR-16 Product Status", "Giữ", "Chỉ Product/Group active cho giao dịch mới."),
    ("BR-17 Warehouse Status", "Sửa", "Kiểm tra cấu hình singleton; người dùng không CRUD/status warehouse."),
    ("BR-18 Location Status", "Giữ", "Không dùng BLOCKED/INACTIVE cho putaway/pick."),
    ("BR-20 Audit Trail", "Giữ", "Audit mọi create/edit/cancel/assign/override/post."),
]
add_table(["BR", "Quyết định", "Nội dung sau chuẩn hóa"], br_rows, [1700, 1600, 6060], font_size=8.5)

add_heading("7.2 Business Rules mới bắt buộc", 2, page_break=True)
new_br = [
    ("BR-WH-01", "Warehouse Singleton", "Hệ thống chỉ có một warehouse cấu hình; mọi nghiệp vụ tự resolve warehouse."),
    ("BR-WH-02", "Location Hierarchy", "Mọi BIN thuộc Zone/Rack hợp lệ; sequence số quyết định sort."),
    ("BR-PR-01", "Group Rotation", "ProductGroup phải có đúng một FIFO/LIFO/FEFO."),
    ("BR-PR-02", "Receipt Layer", "Mỗi lần nhận tạo lớp tồn nội bộ; không gộp các lần nhập chỉ vì thiếu supplier lot."),
    ("BR-PR-03", "FEFO Expiry", "Receipt layer của Product FEFO bắt buộc ExpiryDate hợp lệ."),
    ("BR-QTY-01", "Quantity Precision", "Mọi quantity tuân QuantityScale của UOM."),
    ("BR-IN-01", "Inbound Sources", "Inbound chỉ tham chiếu PO mua về hoặc SO khách trả."),
    ("BR-IN-02", "Receipt before Putaway", "Kiểm nhận và cất hàng là hai phase/state khác nhau."),
    ("BR-IN-03", "Actual Putaway", "Staff xác nhận bin và quantity thực tế; có thể split nhiều bin."),
    ("BR-OUT-01", "Soft/Hard Reservation", "Draft hold có TTL; confirmed SO có hard reservation."),
    ("BR-OUT-02", "Rotation First", "Rotation key đứng trước location sequence."),
    ("BR-OUT-03", "Exact Allocation", "Reservation và PickPlan chỉ rõ Product + Layer + Bin + Quantity."),
    ("BR-OUT-04", "No Early Deduction", "Create/confirm SO và create Outbound không giảm OnHand."),
    ("BR-OUT-05", "Pick to Staging", "Confirm pick chuyển tồn bin sang outbound staging."),
    ("BR-OUT-06", "Final Issue", "Chỉ xác nhận hàng rời kho mới giảm tổng OnHand."),
    ("BR-OUT-07", "Short Pick", "Chỉ post actual; phần thiếu reallocate hoặc review."),
    ("BR-OUT-08", "Atomic Posting", "Balance, reservation, execution, state, ledger và audit là một transaction."),
    ("BR-OUT-09", "Idempotency", "Một command retry không được post hai lần."),
    ("BR-ASG-01", "Exclusive Staff Task", "Chỉ active WAREHOUSE_STAFF không có active inbound/outbound khác được assign."),
    ("BR-RET-01", "Customer Return", "Khách trả tạo lớp tồn mới và phải kiểm định trước AVAILABLE."),
    ("BR-RET-02", "Supplier Return", "Ưu tiên tồn có nguồn từ PO/NCC gốc; xuất khác nguồn cần reason/authorization."),
]
add_table(["ID", "Tên", "Quy tắc"], new_br, [1500, 2300, 5560], font_size=8.4)


add_heading("8. Mô hình dữ liệu đích và migration", 1, page_break=True)
add_heading("8.1 Các entity chính", 2)
add_code(
    "ProductGroup 1 ── n Product\n"
    "Product 1 ── n ReceiptStockLayer\n"
    "ReceiptStockLayer 1 ── n InventoryBalance (mỗi Bin)\n"
    "SalesOrderDetail 1 ── n InventoryReservation\n"
    "OutboundOrderItem 1 ── n OutboundPickPlanLine\n"
    "OutboundPickPlanLine 1 ── n OutboundPickExecution\n"
    "InboundOrderItem 1 ── n InboundReceiptLine\n"
    "InboundReceiptLine 1 ── n InboundPutawayLine\n"
    "InventoryTransaction = immutable ledger cho mọi biến động"
)

add_heading("8.2 Schema change matrix", 2)
schema_rows = [
    ("Warehouses", "Add SingletonKey unique/check hoặc seed invariant; giữ 1 dòng.", "P0"),
    ("StorageLocations", "Add LevelCode, LevelSequence, BinPosition, LocationSequence, IsSystemLocation; remove/deprecate capacity.", "P0"),
    ("ProductGroups", "Add RotationMethod, ShelfLifeDescription.", "P0"),
    ("Products", "Deprecate RotationMethod, TrackLot, TrackExpiry, DefaultShelfLifeDays.", "P1"),
    ("ReceiptStockLayers", "Rename/create từ ProductLots; SystemLayerCode, SupplierLotNumber nullable, ReceivedAt, ExpiryDate, SourceReceiptLineID, Status.", "P0"),
    ("Inventory", "Rename ProductLotID -> ReceiptStockLayerID; add RowVersion; unique Product+Location+Layer.", "P0"),
    ("InventoryReservations", "Layer FK; ReservationType SOFT/HARD; ExpiresAt; Owner references; PARTIALLY_CONSUMED/EXPIRED; RowVersion.", "P0"),
    ("InboundReceiptLines", "Tách dữ liệu kiểm nhận/condition khỏi putaway actual.", "P0"),
    ("InboundPutawayLines", "ReceiptLine, destination bin, planned/actual qty, actor/time.", "P0"),
    ("OutboundPickPlanLines", "Exact source bin/layer/reservation/planned qty/sequence/status.", "P0"),
    ("OutboundPickExecutions", "Actual picked qty, actual source, variance, actor/time, staging transaction.", "P0"),
    ("InventoryTransactions", "Add TransactionGroupID, IdempotencyKey; types RECEIPT, PUTAWAY_OUT/IN, PICK_OUT, STAGE_IN, ISSUE_OUT.", "P0"),
    ("InboundOrders", "Remove ParentInboundOrder; canonical statuses DRAFT/READY/RECEIVING/RECEIVED/PUTAWAY_COMPLETED/CANCELLED.", "P1"),
    ("OutboundOrders", "Canonical statuses DRAFT/READY/PICKING/PICKED/ISSUED/CANCELLED; add RowVersion.", "P0"),
]
add_table(["Bảng", "Thay đổi", "Ưu tiên"], schema_rows, [2300, 5960, 1100], font_size=8.2)

add_heading("8.3 Transaction types", 2, page_break=True)
add_table(
    ["Type", "OnHandDelta", "ReservedDelta", "Location", "Ý nghĩa"],
    [
        ("RECEIPT", "+", "0", "RECEIVING/QUARANTINE", "Ghi nhận hàng vật lý vừa nhận."),
        ("PUTAWAY_OUT", "-", "0", "RECEIVING", "Rời khu nhận."),
        ("PUTAWAY_IN", "+", "0", "BIN", "Vào bin thực tế."),
        ("RESERVE", "0", "+", "BIN", "Giữ mềm/cứng."),
        ("RELEASE_RESERVATION", "0", "-", "BIN", "Giải phóng."),
        ("PICK_OUT", "-", "<=0", "BIN", "Staff lấy thực tế khỏi bin."),
        ("STAGE_IN", "+", "0", "OUTBOUND_STAGING", "Hàng chờ giao."),
        ("ISSUE_OUT", "-", "0", "OUTBOUND_STAGING", "Hàng rời kho."),
        ("STOCKTAKE_ADJUSTMENT", "+/-", "0", "BIN", "Điều chỉnh đã duyệt."),
    ],
    [1800, 1300, 1450, 2200, 2610],
    font_size=8.3,
)
add_body(
    "Các cặp PUTAWAY_OUT/PUTAWAY_IN và PICK_OUT/STAGE_IN dùng chung TransactionGroupID để chứng minh tổng thay đổi toàn kho bằng 0. ISSUE_OUT mới làm tổng kho giảm."
)

add_heading("8.4 Kế hoạch migration dữ liệu", 2)
add_number("Tạo schema mới theo cách additive; chưa xóa cột/bảng cũ.")
add_number("Seed/khóa Warehouse singleton và ba system locations.")
add_number("Sinh LocationSequence từ Zone/Rack/Level/Bin; lập báo cáo mã không parse được để kiểm tra thủ công.")
add_number("Thêm policy ProductGroup. Nếu các Product cùng Group đang có RotationMethod khác nhau, migration phải dừng và yêu cầu tách Group hoặc chọn policy chung.")
add_number("Backfill ProductLots thành ReceiptStockLayers. Lot thật được giữ tại SupplierLotNumber; NO-LOT không được coi là một lớp duy nhất.")
add_number("Dựng lại layer không lô từ InboundOrderDetail + InventoryTransaction theo từng lần nhận. Phân bổ các outbound lịch sử theo rule tại thời điểm giao dịch để tính remaining layer.")
add_number("So sánh tổng OnHand/Reserved trước và sau theo Product+Location; không cut-over khi chênh lệch khác 0.")
add_number("Chạy dual-write tạm thời nếu database đang có người dùng; sau đối soát chuyển read sang schema mới.")
add_number("Xóa/deprecate cột cũ và aliases sau một vòng release ổn định.")
add_callout(
    "Khuyến nghị môi trường đồ án",
    "Nếu dữ liệu hiện tại chủ yếu là seed/test, reseed database theo schema mới an toàn và rẻ hơn việc cố suy ngược tuổi tồn từ NO-LOT đã bị gộp. Nếu là dữ liệu thật, bắt buộc migration có reconciliation report và phê duyệt.",
    "caution",
)


add_heading("9. Kế hoạch thay đổi code và giao diện", 1, page_break=True)
add_heading("9.1 Kiến trúc service đề xuất", 2)
add_code(
    "InventoryAvailabilityService\n"
    "  └─ tính eligible/available stock\n"
    "InventoryAllocationService\n"
    "  ├─ FifoAllocationPolicy\n"
    "  ├─ LifoAllocationPolicy\n"
    "  └─ FefoAllocationPolicy\n"
    "ReservationService\n"
    "  └─ soft hold / hard reserve / release / consume\n"
    "InboundExecutionService\n"
    "  └─ receive / inspect / putaway\n"
    "OutboundExecutionService\n"
    "  └─ build plan / confirm pick / reallocate / confirm issue\n"
    "InventoryPostingService\n"
    "  └─ atomic balance + ledger + idempotency"
)
add_body(
    "Không để SalesOrderService, OutboundOrderService và InventoryRepository tự triển khai ba phiên bản sort khác nhau. Tất cả check availability/reserve/pick plan phải gọi cùng Allocation Engine để một rule cho ra một kết quả."
)

add_heading("9.2 API/command đích", 2)
api_rows = [
    ("GET /api/warehouse", "Xem kho singleton và thống kê vị trí."),
    ("GET/POST/PUT /api/storage-locations/...", "CRUD Zone/Rack/Bin; không có warehouse picker/capacity."),
    ("POST /api/sales-orders/{id}/soft-hold", "Tạo/gia hạn giữ mềm Draft."),
    ("POST /api/sales-orders/{id}/confirm", "Recheck và chuyển hard reservation."),
    ("POST /api/outbound-orders/{id}/build-pick-plan", "Sinh exact allocation atomically."),
    ("POST /api/outbound-orders/{id}/start-picking", "READY -> PICKING; kiểm tra assignee."),
    ("POST /api/outbound-orders/{id}/pick", "Post actual pick BIN -> STAGING; idempotent."),
    ("POST /api/outbound-orders/{id}/reallocate", "Short pick/blocked bin; release+reserve lại."),
    ("POST /api/outbound-orders/{id}/complete-picking", "PICKING -> PICKED khi đủ/đã quyết định exception."),
    ("POST /api/outbound-orders/{id}/issue", "PICKED -> ISSUED; giảm staging và cập nhật SO/PO return."),
    ("POST /api/inbound-orders/{id}/receive", "Kiểm nhận toàn phiếu/dòng; tạo ReceiptStockLayer."),
    ("POST /api/inbound-orders/{id}/putaway", "Xác nhận destination allocations và post RECEIVING -> BIN."),
]
add_table(["Endpoint/Command", "Trách nhiệm"], api_rows, [4200, 5160], font_size=8.5)

add_heading("9.3 UI bắt buộc", 2)
add_heading("Warehouse/Location", 3)
for text in [
    "Trang Warehouse Overview read-only; không nút Thêm/Sửa/Xóa kho.",
    "Location Master dạng cây Zone/Rack/Bin; CRUD ở node; hiển thị Level/Bin sequence và trạng thái.",
    "Không hiển thị diện tích, tải trọng, thể tích hoặc occupancy theo capacity.",
]:
    add_bullet(text)
add_heading("Product Group/Product", 3)
for text in [
    "Màn cấu hình Group có select bắt buộc FIFO/LIFO/FEFO và text HSD.",
    "Bảng Group không thêm hai cột này; detail/config mới hiển thị.",
    "Product create/edit bỏ TrackLot/TrackExpiry/Rotation riêng; hiển thị policy hiệu lực từ Group.",
]:
    add_bullet(text)
add_heading("Inbound", 3)
for text in [
    "Create không có lựa chọn nhập bổ sung; source selector chỉ PO hoặc SO Return.",
    "Receive validate toàn bộ dòng; số lô không bắt buộc; ExpiryDate xuất hiện/bắt buộc khi FEFO.",
    "Putaway có search/filter và breadcrumb Zone/Rack/Level/Bin; split quantity; còn chờ cất phải hiển thị rõ.",
]:
    add_bullet(text)
add_heading("Outbound", 3)
for text in [
    "Create không hỏi kho nguồn; hệ thống tự dùng kho singleton.",
    "Process hiển thị PickPlan theo sequence, ngày nhập/HSD, vị trí, planned/actual và exception.",
    "Nút 'Lưu lần lấy' không được ghi nhận là đã xuất; phải có 'Hoàn tất lấy hàng' và 'Xác nhận hàng rời kho'.",
    "Hiển thị riêng OnHand, Reserved, Available, Picked/Staged và Issued; không dùng một cột Số lượng chung.",
]:
    add_bullet(text)

add_heading("9.4 Thứ tự triển khai", 2)
phases = [
    ("P0.1", "Schema additive + seed singleton/system locations", "Không làm thay đổi tồn hiện tại."),
    ("P0.2", "ReceiptStockLayer + migration/reseed", "Loại NO-LOT merge; giữ tổng tồn."),
    ("P0.3", "Allocation Engine + three policies", "Dùng chung cho reserve và pick plan."),
    ("P0.4", "PickPlan + reservation concurrency", "Exact source/quantity; row version/idempotency."),
    ("P0.5", "Split pick/staging/final issue", "Đảm bảo posting points và state machine."),
    ("P1.1", "ProductGroup policy + Product UI", "Chuyển source of truth sang group."),
    ("P1.2", "Warehouse/Location UI cleanup", "Bỏ CRUD kho/capacity; thêm sequence."),
    ("P1.3", "Inbound state/receipt/putaway cleanup", "Bỏ supplemental; canonical status."),
    ("P1.4", "Reports/audit/alerts", "Expiry by layer; staged stock; exception history."),
    ("P1.5", "Docs + regression + UAT", "Cập nhật đồng bộ trước nghiệm thu."),
]
add_table(["Phase", "Hạng mục", "Điều kiện hoàn thành"], phases, [1200, 4200, 3960], font_size=8.6)


add_heading("10. Ma trận cập nhật tài liệu đặc tả", 1, page_break=True)
doc_rows = [
    ("Business Rule.pdf", "BR-01/02/03/04/07/09/10/17; thêm LIFO, receipt layer, staging, idempotency.", "Sửa"),
    ("Nghiệp vụ.pdf", "Thêm LIFO; rotation-first; bỏ capacity; reservation draft/confirmed; pick-to-stage.", "Sửa"),
    ("Product Group config.pdf", "Bỏ quản lý lô bắt buộc; policy nằm Group; FEFO yêu cầu expiry per receipt layer.", "Sửa"),
    ("(New) 63 UC.pdf", "UC18 bỏ supplemental; UC22 bỏ bắt lot; UC23 Staff chọn actual bin; UC31 tách pick/final issue; UC48 đổi tên.", "Sửa"),
    ("Report 3.1 UCs", "UC01-09, 17-32, 48, 64-65; glossary; request fields; AF; verification criteria; state machines.", "Sửa lớn"),
    ("BP-01", "Thể hiện RECEIVING/QUARANTINE, ReceiptStockLayer, putaway actual và inventory posting.", "Sửa"),
    ("BP-02", "Customer return tạo layer mới, quality hold và putaway.", "Sửa"),
    ("BP-03", "Soft/hard reserve; FIFO/LIFO/FEFO; PickPlan; staging; final issue.", "Sửa lớn"),
    ("BP-04", "Bỏ capacity validation; chỉ transfer nội bộ độc lập.", "Sửa nhẹ"),
    ("BP-06", "PO return reserve đúng nguồn, pick/stage/issue.", "Sửa"),
    ("Sequence Diagrams", "UC22/23/26/30/31 và UC48; thêm Allocation/Reservation/Posting services và hai command pick/issue.", "Sửa lớn"),
    ("ERD.pdf", "Singleton warehouse; location sequence; ProductGroup policy; ReceiptStockLayer; plan/execution; staging ledger.", "Sửa lớn"),
    ("Feature list", "Giữ F05.03 nhưng mô tả ba phase Pick - Stage - Issue; F04.03 actual putaway.", "Sửa"),
]
add_table(["Tài liệu", "Nội dung cần thay đổi", "Mức"], doc_rows, [2300, 5960, 1100], font_size=8.3)

add_heading("10.1 State machine cần thay trong Report 3.1", 2)
add_table(
    ["Entity", "Hiện tại", "Đích"],
    [
        ("Product Lot", "ACTIVE/EXPIRED/DEPLETED/BLOCKED", "ReceiptStockLayer: AVAILABLE/QUARANTINED/EXPIRED/DEPLETED/BLOCKED."),
        ("Inbound", "DRAFT -> READY -> RECEIVING -> RECEIVED -> PUTAWAY_COMPLETED", "Giữ state logic; database dùng trực tiếp, không alias ASSIGNED/COMPLETED."),
        ("Outbound", "DRAFT/RESERVED -> READY -> ISSUING -> ISSUED", "DRAFT -> READY -> PICKING -> PICKED -> ISSUED; reservation là entity riêng."),
        ("Reservation", "ACTIVE -> CONSUMED/RELEASED/EXPIRED", "Thêm SOFT/HARD type và PARTIALLY_CONSUMED; TTL cho SOFT."),
        ("SO", "DRAFT -> CONFIRMED -> PARTIALLY_ISSUED -> ISSUED", "DRAFT có soft hold; CONFIRMED có hard reserve; fulfillment chỉ tăng ở final issue."),
    ],
    [1900, 3400, 4060],
)

add_heading("10.2 UC30 và UC31 phải viết lại", 2)
add_body(
    "UC30 không còn 'nearest location first, then FEFO/FIFO'. Câu đúng là: lọc eligible stock; áp dụng ProductGroup RotationMethod; trong cùng priority group áp dụng LocationSequence; atomically create reservation và PickPlan."
)
add_body(
    "UC31 không còn một thao tác vừa pick vừa issue. Main flow phải có xác nhận actual pick vào staging, short-pick/reallocation và xác nhận final issue. Các alternative flow phải bổ sung idempotent retry, expired FEFO layer, stocktake lock, cancellation after pick và return-to-bin."
)


add_heading("11. Kiểm thử, nghiệm thu và triển khai", 1, page_break=True)
add_heading("11.1 Ma trận test bắt buộc", 2)
tests = [
    ("FIFO-01", "Hai ngày nhập, nhiều bin", "Pick ngày cũ; cùng ngày A->Z, bin nhỏ->lớn."),
    ("LIFO-01", "Hai ngày nhập, nhiều bin", "Pick ngày mới; cùng ngày Z->A, bin lớn->nhỏ."),
    ("FEFO-01", "Ngày nhập mới nhưng HSD sớm", "HSD sớm được pick trước bất kể location."),
    ("FEFO-02", "Thiếu HSD", "Không cấp phát; layer bị HOLD/MISSING_EXPIRY."),
    ("QTY-01", "bao/cây QuantityScale=0", "Reject 1.5."),
    ("QTY-02", "tấn QuantityScale=3", "Accept 1.250; reject 1.2501."),
    ("RES-01", "Hai Sales đồng thời", "Chỉ transaction đầu giữ đủ; không âm/over-reserve."),
    ("RES-02", "Draft hold hết TTL", "Reservation EXPIRED và Available được trả lại."),
    ("PICK-01", "Pick một line nhiều bin", "Plan và actual đúng tổng; bin->staging bằng group transaction."),
    ("PICK-02", "Nhấn đúp xác nhận", "Idempotency trả cùng kết quả; không trừ hai lần."),
    ("PICK-03", "Short pick", "Chỉ trừ actual; reallocate phần còn thiếu."),
    ("ISSUE-01", "Đã pick chưa issue", "Tổng kho không giảm; staging có hàng; SO chưa fulfilled."),
    ("ISSUE-02", "Final issue", "Staging giảm; tổng kho và SO fulfilled cập nhật atomically."),
    ("CANCEL-01", "Cancel READY", "Plan hủy; reservation theo owner policy."),
    ("CANCEL-02", "Cancel sau pick", "Không hủy trực tiếp; bắt return-to-bin."),
    ("IN-01", "Receipt FEFO", "Không cần supplier lot; expiry bắt buộc; layer tự sinh."),
    ("IN-02", "Putaway split", "Nhiều bin; tổng không vượt accepted; receiving giảm đúng."),
    ("RET-01", "Khách trả", "Layer mới, quarantine trước available."),
    ("RET-02", "Trả NCC", "Ưu tiên layer từ PO gốc; override có reason."),
    ("ASG-01", "Staff đang có task", "Không xuất hiện trong selector và API reject race condition."),
]
add_table(["ID", "Kịch bản", "Kết quả mong đợi"], tests, [1300, 3400, 4660], font_size=8.2)

add_heading("11.2 Acceptance criteria tổng", 2)
for text in [
    "Tổng OnHand trước/sau migration khớp theo Product và toàn kho; không có Available âm.",
    "Một tập dữ liệu cố định luôn sinh cùng PickPlan sequence.",
    "Không tồn tại đường code nào sort FEFO sau LocationCode hoặc xử lý LIFO bằng nhánh FIFO.",
    "Không còn request field LotNumber required cho sản phẩm không cần supplier lot; FEFO thiếu ExpiryDate bị chặn đúng dòng.",
    "Nhập, putaway, reserve, pick, staging và issue đều có immutable ledger/audit và người thực hiện.",
    "Màn Staff chỉ cho thao tác phiếu được assign; một Staff không có hai active tasks kể cả khi hai request assign đồng thời.",
    "Report 3.1, BP, Sequence, ERD, Business Rules và UI dùng cùng thuật ngữ/state.",
]:
    add_bullet(text)

add_heading("11.3 Out of scope có chủ đích", 2)
for text in [
    "Tính sức chứa bin theo trọng lượng, thể tích hoặc diện tích.",
    "Tối ưu đường đi ngắn nhất bằng tọa độ, bản đồ hoặc thuật toán routing.",
    "RFID/serial cho từng bao, cây hoặc cuộn.",
    "Truy xuất/thu hồi theo supplier manufacturing lot khi supplier không cung cấp mã.",
    "Tài chính, đơn giá, hóa đơn, giá vốn, công nợ và chứng từ kế toán.",
    "Điều phối vận tải/chuyến xe và proof-of-delivery bên ngoài kho.",
    "Quy đổi nhiều UOM nếu chưa có bảng conversion được phê duyệt.",
]:
    add_bullet(text)

add_heading("11.4 Rủi ro chính và kiểm soát", 2)
add_table(
    ["Rủi ro", "Mức", "Kiểm soát"],
    [
        ("NO-LOT lịch sử không tách được tuổi tồn còn lại", "Cao", "Reconstruct từ ledger hoặc reseed; reconciliation bắt buộc."),
        ("Đổi ProductGroup policy khi đang có tồn", "Cao", "Chặn đổi hoặc chạy reallocation validation/migration có phê duyệt."),
        ("Status alias gây dữ liệu không nhất quán", "Cao", "Canonical enum + transition service + DB CHECK."),
        ("Double-click/retry trừ tồn hai lần", "Cao", "IdempotencyKey unique + transaction."),
        ("Staff chọn sai bin/layer", "Trung bình", "PickPlan exact, scan/mã vị trí, override reason/approval."),
        ("Draft giữ hàng vô thời hạn", "Trung bình", "TTL + cleanup job + hiển thị thời gian hết hạn."),
        ("Không capacity dẫn đến Staff chọn bin không đủ chỗ", "Chấp nhận", "Hệ thống chỉ đề xuất; Staff chịu trách nhiệm xác nhận thực tế."),
    ],
    [3400, 1300, 4660],
    font_size=8.4,
)


add_heading("Phụ lục A. Trường dữ liệu tối thiểu", 1, page_break=True)
add_heading("A.1 ReceiptStockLayer", 2)
add_table(
    ["Field", "Rule"],
    [
        ("ReceiptStockLayerId", "PK; hệ thống."),
        ("SystemLayerCode", "Unique; hệ thống tự sinh."),
        ("ProductId", "FK Product; bắt buộc."),
        ("InboundReceiptLineId", "Nguồn nhận; bắt buộc cho inbound mới."),
        ("ReceivedAt", "UTC datetime; bắt buộc."),
        ("ReceivedBusinessDate", "Ngày nghiệp vụ theo Asia/Saigon; dùng nhóm 'cùng ngày'."),
        ("ExpiryDate", "Bắt buộc khi effective RotationMethod=FEFO."),
        ("SupplierLotNumber", "Nullable; không validation bắt buộc."),
        ("Status", "AVAILABLE, QUARANTINED, BLOCKED, EXPIRED, DEPLETED."),
    ],
    [3300, 6060],
)

add_heading("A.2 PickPlanLine", 2)
add_table(
    ["Field", "Rule"],
    [
        ("OutboundPickPlanLineId", "PK."),
        ("OutboundOrderItemId", "FK; một outbound item có nhiều plan lines."),
        ("InventoryReservationId", "FK nullable cho plan chưa hard-reserved; bắt buộc trước READY."),
        ("ReceiptStockLayerId", "Exact layer."),
        ("SourceLocationId", "Exact source BIN."),
        ("PlannedQuantity", ">0, đúng QuantityScale."),
        ("SequenceNo", "Số thứ tự deterministic."),
        ("RotationMethodSnapshot", "FIFO/LIFO/FEFO tại thời điểm plan; phục vụ audit."),
        ("Status", "PLANNED, PARTIALLY_PICKED, PICKED, SKIPPED, CANCELLED."),
    ],
    [3300, 6060],
)

add_heading("A.3 InventoryReservation", 2)
add_table(
    ["Field", "Rule"],
    [
        ("ReservationType", "SOFT hoặc HARD."),
        ("Owner", "SalesOrderDetail hoặc OutboundOrderItem phù hợp source."),
        ("Product/Layer/Location", "Exact allocation key."),
        ("ReservedQuantity", ">0."),
        ("ConsumedQuantity", "0..ReservedQuantity."),
        ("Status", "ACTIVE, PARTIALLY_CONSUMED, CONSUMED, RELEASED, EXPIRED."),
        ("ExpiresAt", "Bắt buộc với SOFT; nullable với HARD."),
        ("RowVersion", "Concurrency token."),
    ],
    [3300, 6060],
)


add_heading("Phụ lục B. Thuật ngữ chuẩn", 1, page_break=True)
glossary = [
    ("Kho", "Warehouse singleton trong BMWMS."),
    ("Khu", "Zone, cấp cao nhất của sơ đồ vị trí trong kho."),
    ("Rack", "Khung/kệ thuộc một Zone."),
    ("Tầng", "Level A-Z nằm trong Rack; lưu như thuộc tính của Bin."),
    ("Bin", "Ô/vị trí chi tiết đặt hàng; có Position 1-n."),
    ("Lô NCC", "Mã do nhà sản xuất/nhà cung cấp cấp; tùy chọn."),
    ("Lớp tồn theo lần nhập", "Nhóm số lượng nội bộ có chung lần nhận, ReceivedAt và ExpiryDate."),
    ("OnHand", "Số lượng vật lý tại một location."),
    ("Reserved", "Phần OnHand đã giữ, chưa được cấp phát lại."),
    ("Available", "Phần đủ điều kiện cấp phát mới."),
    ("Putaway", "Cất hàng đã nhận từ RECEIVING vào Bin thực tế."),
    ("Pick", "Lấy hàng từ Bin theo PickPlan."),
    ("Staging", "Vị trí chờ giao; hàng đã rời Bin nhưng chưa rời kho."),
    ("Issue", "Xác nhận hàng rời kho và giảm tổng OnHand."),
    ("Reservation", "Giữ quyền sử dụng quantity; không phải trừ tồn."),
    ("Allocation", "Thuật toán chọn Layer + Bin + Quantity theo rotation rule."),
]
add_table(["Thuật ngữ", "Định nghĩa"], glossary, [3000, 6360], font_size=8.7)


add_heading("Phụ lục C. Dẫn chiếu hiện trạng code", 1, page_break=True)
refs = [
    ("Database_v3.sql:758-860", "Warehouse/Zone/Rack/Location và các capacity fields."),
    ("Database_v3.sql:888-1030", "ProductGroup thiếu policy; Product có Rotation/TrackLot; ProductLots bắt LotNumber."),
    ("Database_v3.sql:1123-1325", "Inbound/Outbound statuses và ProductLot FK."),
    ("Database_v3.sql:1458-1535", "Inventory và immutable transaction structure."),
    ("BMWMS.Business/Services/WarehouseService.cs", "Create/Update/Delete multi-warehouse cần bỏ."),
    ("BMWMS.Business/Services/ProductGroupService.cs", "DTO/service group chưa có RotationMethod/ShelfLifeDescription."),
    ("BMWMS.Business/Services/InboundService.cs:1193-1218", "TrackLot validation và NO-LOT reuse."),
    ("BMWMS.Business/Services/InboundService.cs:1308-1410", "Putaway Staff chọn vị trí và post INBOUND - phần nên giữ/refactor."),
    ("BMWMS.Repository/Repositories/Inventory/InventoryRepository.cs:184-255", "Reservation FEFO/else FIFO; thiếu LIFO."),
    ("BMWMS.Business/Services/Inventory/OutboundOrderService.cs:371-503", "ExecutePick giảm OnHand và complete ngay."),
    ("BMWMS.Business/Services/Inventory/OutboundOrderService.cs:512-580", "Allocation order theo LocationCode; FEFO ordering sai ưu tiên."),
    ("docs/schema_patch_separate_order_tasks_from_transfers.sql", "Quyết định tách Inbound/Outbound khỏi TransferOrder đã đúng hướng."),
]
add_table(["Nguồn", "Nhận xét"], refs, [4100, 5260], font_size=8.4)


add_heading("Phụ lục D. Nguồn tham khảo nghiệp vụ", 1, page_break=True)
sources = [
    ("GS1 Global Traceability Standard", "Batch/lot phân biệt nhóm vật phẩm có chung đặc điểm truy vết; mức truy vết phải phù hợp mục tiêu và chi phí.", "https://www.gs1.org/standards/gs1-global-traceability-standard/current-standard"),
    ("Microsoft - Inventory picking aging", "FIFO/LIFO có thể dựa trên ngày hàng vào kho ngay cả với tồn không quản lý batch.", "https://learn.microsoft.com/en-us/dynamics365/supply-chain/warehousing/location-directive-inventory-picking-aging"),
    ("Microsoft - Item tracing", "Batch/serial/vendor batch hỗ trợ truy ngược nguồn và truy xuôi nơi đã xuất.", "https://learn.microsoft.com/en-us/dynamics365/supply-chain/inventory/trace-items-raw-materials-inventory-production-sales"),
    ("Oracle WMS - Allocation methods", "FIFO theo create timestamp ASC, LIFO DESC, FEFO theo expiry/priority date ASC; location pick sequence là secondary ordering.", "https://docs.oracle.com/en/cloud/saas/warehouse-management/26a/owmol/optional-step-additional-configuration-parameters.html"),
]
add_table(["Nguồn", "Điểm dùng", "Đường dẫn"], sources, [2400, 3760, 3200], font_size=7.9)

add_heading("Kết luận phê duyệt", 1)
add_callout(
    "Baseline đề xuất",
    "Triển khai theo ReceiptStockLayer + Rotation-first Allocation + Soft/Hard Reservation + PickPlan + OUTBOUND_STAGING. Đây là thay đổi cấu trúc dữ liệu và state machine, không phải chỉnh giao diện cục bộ. Chỉ bắt đầu code sau khi nhóm chấp thuận các quyết định D-01 đến D-15 và kế hoạch migration dữ liệu.",
    "decision",
)

# Document settings: update fields on open and language metadata.
settings = doc.settings._element
update_fields = settings.find(qn("w:updateFields"))
if update_fields is None:
    update_fields = OxmlElement("w:updateFields")
    settings.append(update_fields)
update_fields.set(qn("w:val"), "true")

core = doc.core_properties
core.title = "BMWMS - Đặc tả nghiệp vụ chốt Kho, Product, Inbound và Outbound"
core.subject = "Business rules, target data model and implementation plan"
core.author = "BMWMS Project Team"
core.keywords = "BMWMS, warehouse, FIFO, LIFO, FEFO, inbound, outbound, reservation, picking"
core.comments = "Generated as a consolidated decision baseline for branch trungdq."

OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT_DOCX)
print(OUT_DOCX)
