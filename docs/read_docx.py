from docx import Document

doc = Document(r'c:\Users\Dumb Trung\Documents\GitHub\bmwms\docs\Functional Specification.docx')
text = []
for para in doc.paragraphs:
    if para.text.strip():
        text.append(para.text)

for table in doc.tables:
    for row in table.rows:
        row_text = []
        for cell in row.cells:
            row_text.append(cell.text.replace('\n', ' ').strip())
        text.append(' | '.join(row_text))

with open(r'c:\Users\Dumb Trung\Documents\GitHub\bmwms\docs\func_spec_python.txt', 'w', encoding='utf-8') as f:
    f.write('\n'.join(text))
